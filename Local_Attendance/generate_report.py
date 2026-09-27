"""
ESSL X990 Local Attendance Report Generator (v6)
- Daily Attendance & Monthly Summary: S.No + Name (A-Z), no User ID column
- Raw Punches: sorted by timestamp (User ID retained for audit)
"""

from zk import ZK
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, time
from collections import defaultdict
import os
import sys

# ----------- CONFIGURATION -----------
DEVICE_IP = "192.168.1.34"
DEVICE_PORT = 4370
TIMEOUT = 10
OUTPUT_DIR = "reports"
STANDARD_WORKDAY_HOURS = 9
NOON = time(12, 0, 0)
DUPLICATE_WINDOW_SECONDS = 60
MAX_REASONABLE_HOURS = 14
MIN_REASONABLE_HOURS = 1
# --------------------------------------


def connect_device():
    print(f"Connecting to ESSL device at {DEVICE_IP}:{DEVICE_PORT} ...")
    zk = ZK(DEVICE_IP, port=DEVICE_PORT, timeout=TIMEOUT,
            password=0, force_udp=False, ommit_ping=False)
    try:
        conn = zk.connect()
        print("Connected successfully.")
        return conn
    except Exception as e:
        print(f"Connection failed: {e}")
        sys.exit(1)


def fetch_data(conn):
    users = conn.get_users()
    user_map = {u.user_id: u.name for u in users}
    attendance = conn.get_attendance()
    print(f"Retrieved {len(attendance)} attendance records.")
    return attendance, user_map


def detect_duplicates(timestamps):
    duplicates = []
    sorted_ts = sorted(timestamps)
    for i in range(1, len(sorted_ts)):
        delta = (sorted_ts[i] - sorted_ts[i - 1]).total_seconds()
        if delta <= DUPLICATE_WINDOW_SECONDS:
            duplicates.append(sorted_ts[i])
    return duplicates


def process_attendance(attendance, user_map):
    grouped = defaultdict(list)
    for record in attendance:
        key = (record.user_id, record.timestamp.date())
        grouped[key].append(record)

    all_punch_times_by_second = defaultdict(list)
    for record in attendance:
        sec_key = record.timestamp.replace(microsecond=0)
        all_punch_times_by_second[sec_key].append(record.user_id)

    daily_records = []
    for (user_id, date), records in grouped.items():
        timestamps = sorted([r.timestamp for r in records])
        morning = [t for t in timestamps if t.time() < NOON]
        afternoon = [t for t in timestamps if t.time() >= NOON]

        clock_in = min(morning) if morning else None
        clock_out = max(afternoon) if afternoon else None

        total_hours = None
        if clock_in and clock_out:
            total_hours = round((clock_out - clock_in).total_seconds() / 3600, 2)

        excess = deficit = None
        if total_hours is not None:
            diff = round(total_hours - STANDARD_WORKDAY_HOURS, 2)
            excess = diff if diff > 0 else 0
            deficit = abs(diff) if diff < 0 else 0

        flags = []
        if detect_duplicates(timestamps):
            flags.append(f"Duplicate punch(es) within {DUPLICATE_WINDOW_SECONDS}s")
        if len(timestamps) == 1:
            flags.append("Only 1 punch recorded (missing pair)")
        elif len(timestamps) > 4:
            flags.append(f"Unusually high punch count ({len(timestamps)})")
        if total_hours is not None:
            if total_hours < MIN_REASONABLE_HOURS:
                flags.append(f"Suspiciously short day ({total_hours}h)")
            elif total_hours > MAX_REASONABLE_HOURS:
                flags.append(f"Suspiciously long day ({total_hours}h)")
        if clock_in and clock_out and clock_in == clock_out:
            flags.append("Clock-in and clock-out identical")
        if date.weekday() >= 5:
            flags.append("Weekend activity")
        for t in timestamps:
            sec_key = t.replace(microsecond=0)
            other_users = set(all_punch_times_by_second[sec_key]) - {user_id}
            if other_users:
                flags.append(f"Same-second punch as user(s) {sorted(other_users)} — possible buddy punching")
                break

        daily_records.append({
            "user_id": user_id,
            "name": user_map.get(user_id, "Unknown"),
            "date": date,
            "clock_in": clock_in,
            "clock_out": clock_out,
            "total_hours": total_hours,
            "excess_hours": excess,
            "deficit_hours": deficit,
            "raw_punches": [t.strftime("%H:%M:%S") for t in timestamps],
            "flags": flags
        })

    daily_records.sort(key=lambda r: (r["name"].lower(), r["date"]))
    return daily_records


def build_monthly_summary(daily_records):
    summary = defaultdict(lambda: {"name": "", "days_present": 0, "total_hours": 0.0,
                                    "excess_hours": 0.0, "deficit_hours": 0.0, "flagged_days": 0})
    for rec in daily_records:
        uid = rec["user_id"]
        s = summary[uid]
        s["name"] = rec["name"]
        s["days_present"] += 1
        if rec["total_hours"]:
            s["total_hours"] += rec["total_hours"]
        if rec["excess_hours"]:
            s["excess_hours"] += rec["excess_hours"]
        if rec["deficit_hours"]:
            s["deficit_hours"] += rec["deficit_hours"]
        if rec["flags"]:
            s["flagged_days"] += 1

    rows = []
    for uid, data in summary.items():
        rows.append({
            "name": data["name"],
            "days_present": data["days_present"],
            "total_hours": round(data["total_hours"], 2),
            "excess_hours": round(data["excess_hours"], 2),
            "deficit_hours": round(data["deficit_hours"], 2),
            "equivalent_full_days": round(data["total_hours"] / STANDARD_WORKDAY_HOURS, 2),
            "flagged_days": data["flagged_days"]
        })
    rows.sort(key=lambda r: r["name"].lower())
    return rows


def style_header(ws, ncols):
    fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def autofit_columns(ws):
    for col in ws.columns:
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4


def export_to_excel(daily_records, summary_rows, raw_attendance, user_map):
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    ts = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    filename = os.path.join(OUTPUT_DIR, f"attendance_report_{ts}.xlsx")
    wb = Workbook()

    # ---------- Sheet 1: Daily Attendance ----------
    ws1 = wb.active
    ws1.title = "Daily Attendance"
    headers1 = ["S.No", "Name", "Date", "Clock-In", "Clock-Out",
                "Total Hours", "Excess Hours", "Deficit Hours", "Flags"]
    ws1.append(headers1)
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    for idx, rec in enumerate(daily_records, start=1):
        row = [
            idx, rec["name"], rec["date"].strftime("%Y-%m-%d"),
            rec["clock_in"].strftime("%H:%M:%S") if rec["clock_in"] else "Missing",
            rec["clock_out"].strftime("%H:%M:%S") if rec["clock_out"] else "Missing",
            rec["total_hours"] if rec["total_hours"] is not None else "N/A",
            rec["excess_hours"] if rec["excess_hours"] else 0,
            rec["deficit_hours"] if rec["deficit_hours"] else 0,
            "; ".join(rec["flags"]) if rec["flags"] else ""
        ]
        ws1.append(row)
        if rec["flags"]:
            for cell in ws1[ws1.max_row]:
                cell.fill = red_fill
    style_header(ws1, len(headers1))
    autofit_columns(ws1)
    ws1.freeze_panes = "A2"

    # ---------- Sheet 2: Monthly Summary ----------
    ws2 = wb.create_sheet("Monthly Summary")
    headers2 = ["S.No", "Name", "Days Present", "Total Hours",
                "Excess Hours", "Deficit Hours",
                f"Equivalent Full Days ({STANDARD_WORKDAY_HOURS}h)", "Flagged Days"]
    ws2.append(headers2)
    for idx, row in enumerate(summary_rows, start=1):
        ws2.append([idx, row["name"], row["days_present"], row["total_hours"],
                     row["excess_hours"], row["deficit_hours"],
                     row["equivalent_full_days"], row["flagged_days"]])
    style_header(ws2, len(headers2))
    autofit_columns(ws2)
    ws2.freeze_panes = "A2"

    # ---------- Sheet 3: Anomalies ----------
    ws3 = wb.create_sheet("Anomalies")
    headers3 = ["S.No", "Name", "Date", "Raw Punches", "Flags"]
    ws3.append(headers3)

    flagged = [rec for rec in daily_records if rec["flags"]]
    flagged.sort(key=lambda r: (r["name"].lower(), r["date"]))
    for idx, rec in enumerate(flagged, start=1):
        ws3.append([idx, rec["name"], rec["date"].strftime("%Y-%m-%d"),
                    ", ".join(rec["raw_punches"]), "; ".join(rec["flags"])])
    style_header(ws3, len(headers3))
    autofit_columns(ws3)
    ws3.freeze_panes = "A2"

    # ---------- Sheet 4: Raw Punches ----------
    ws4 = wb.create_sheet("Raw Punches")
    headers4 = ["S.No", "Timestamp", "Date", "Time", "User ID", "Name",
                "Status Code", "Status Label", "Punch", "UID"]
    ws4.append(headers4)

    status_map = {0: "Check-In", 1: "Check-Out", 2: "Break-Out",
                  3: "Break-In", 4: "OT-In", 5: "OT-Out"}

    raw_sorted = sorted(raw_attendance, key=lambda r: r.timestamp)
    for idx, r in enumerate(raw_sorted, start=1):
        ws4.append([
            idx,
            r.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
            r.timestamp.strftime("%Y-%m-%d"),
            r.timestamp.strftime("%H:%M:%S"),
            r.user_id,
            user_map.get(r.user_id, "Unknown"),
            getattr(r, "status", None),
            status_map.get(getattr(r, "status", None), "Unknown"),
            getattr(r, "punch", None),
            getattr(r, "uid", None)
        ])
    style_header(ws4, len(headers4))
    autofit_columns(ws4)
    ws4.freeze_panes = "A2"

    wb.save(filename)
    print(f"Report saved: {os.path.abspath(filename)}")
    return filename


def main():
    conn = connect_device()
    try:
        attendance, user_map = fetch_data(conn)
        if not attendance:
            print("No attendance records found.")
            return
        daily_records = process_attendance(attendance, user_map)
        summary_rows = build_monthly_summary(daily_records)
        export_to_excel(daily_records, summary_rows, attendance, user_map)
    finally:
        conn.disconnect()
        print("Disconnected.")


if __name__ == "__main__":
    main()
