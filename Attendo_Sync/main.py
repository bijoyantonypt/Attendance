"""
AttendanceTracker — main GUI entry point.
On launch: shows Device IP field, auto-fetches attendance data,
updates both databases, computes payroll, and displays the dashboard.
"""

import os
import sys
import json
import threading
import sqlite3                                          # ── EXPORT ADDITION ──
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
from openpyxl import Workbook                           # ── EXPORT ADDITION ──
from openpyxl.styles import Font, PatternFill, Alignment  # ── EXPORT ADDITION ──
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg

import attendance_engine
import db_manager
import payroll_manager
import charts


def get_base_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


BASE_DIR = get_base_dir()
CONFIG_PATH = os.path.join(BASE_DIR, "config.json")


def load_config():
    default = {
        "device_ip": "192.168.29.201",
        "drive_sync_folder": "",
        "default_hourly_rate": 100.0,
        "standard_workday_hours": 9,
    }
    if os.path.exists(CONFIG_PATH):
        with open(CONFIG_PATH, "r") as f:
            default.update(json.load(f))
    return default


def save_config(cfg):
    with open(CONFIG_PATH, "w") as f:
        json.dump(cfg, f, indent=2)


# ── EXPORT ADDITION ── helper functions mirroring Local_Attendance styling ──

def _style_header(ws, ncols):
    fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def _autofit_columns(ws):
    for col in ws.columns:
        max_len = max(
            (len(str(c.value)) if c.value is not None else 0) for c in col
        )
        ws.column_dimensions[col[0].column_letter].width = max_len + 4


def export_attendance_to_excel(attendance_db_path, output_path, standard_workday_hours=9):
    """
    Reads daily_attendance from SQLite and writes a styled 3-sheet Excel file:
      Sheet 1 – Daily Attendance
      Sheet 2 – Monthly Summary
      Sheet 3 – Raw Punches (date + clock_in + clock_out per row as stored)
    Mirrors the export format of Local_Attendance/generate_report.py exactly.
    """
    if not os.path.exists(attendance_db_path):
        raise FileNotFoundError("attendance.db not found. Fetch data first.")

    conn = sqlite3.connect(attendance_db_path)
    try:
        import pandas as pd
        daily_df = pd.read_sql(
            "SELECT * FROM daily_attendance ORDER BY name, date", conn
        )
        monthly_df = pd.read_sql(
            "SELECT * FROM monthly_summary ORDER BY name, year_month", conn
        )
    finally:
        conn.close()

    wb = Workbook()
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    # ── Sheet 1: Daily Attendance ──
    ws1 = wb.active
    ws1.title = "Daily Attendance"
    headers1 = [
        "S.No", "Name", "Date", "Clock-In", "Clock-Out",
        "Total Hours", "Excess Hours", "Deficit Hours", "Flags",
    ]
    ws1.append(headers1)

    for idx, row in enumerate(daily_df.itertuples(index=False), start=1):
        data_row = [
            idx,
            row.name,
            row.date,
            row.clock_in if row.clock_in else "Missing",
            row.clock_out if row.clock_out else "Missing",
            row.total_hours if row.total_hours is not None else "N/A",
            row.excess_hours if row.excess_hours else 0,
            row.deficit_hours if row.deficit_hours else 0,
            row.flags if row.flags else "",
        ]
        ws1.append(data_row)
        if row.flags:                          # red highlight for flagged rows
            for cell in ws1[ws1.max_row]:
                cell.fill = red_fill

    _style_header(ws1, len(headers1))
    _autofit_columns(ws1)
    ws1.freeze_panes = "A2"

    # ── Sheet 2: Monthly Summary ──
    ws2 = wb.create_sheet("Monthly Summary")
    headers2 = [
        "S.No", "Name", "Month", "Days Present",
        "Total Hours", "Excess Hours", "Deficit Hours",
        f"Equivalent Full Days ({standard_workday_hours}h)", "Flagged Days",
    ]
    ws2.append(headers2)

    for idx, row in enumerate(monthly_df.itertuples(index=False), start=1):
        ws2.append([
            idx,
            row.name,
            row.year_month,
            row.days_present,
            row.total_hours,
            row.excess_hours,
            row.deficit_hours,
            row.equivalent_full_days,
            row.flagged_days,
        ])

    _style_header(ws2, len(headers2))
    _autofit_columns(ws2)
    ws2.freeze_panes = "A2"

    # ── Sheet 3: Raw Punches (all daily records as audit log) ──
    ws3 = wb.create_sheet("Raw Punches")
    headers3 = ["S.No", "Name", "Date", "Clock-In", "Clock-Out", "Flags"]
    ws3.append(headers3)

    for idx, row in enumerate(daily_df.itertuples(index=False), start=1):
        ws3.append([
            idx,
            row.name,
            row.date,
            row.clock_in if row.clock_in else "—",
            row.clock_out if row.clock_out else "—",
            row.flags if row.flags else "",
        ])

    _style_header(ws3, len(headers3))
    _autofit_columns(ws3)
    ws3.freeze_panes = "A2"

    wb.save(output_path)

# ── END EXPORT ADDITION ──────────────────────────────────────────────────────


class AttendanceApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Attendance Tracker & Payroll Dashboard")
        self.root.geometry("1100x750")
        self.cfg = load_config()
        self.attendance_db = db_manager.get_db_path(BASE_DIR)
        self.payroll_db = payroll_manager.get_db_path(BASE_DIR)
        self._build_top_bar()
        self._build_tabs()
        self.status_var.set("Ready.")
        self.root.after(500, self.fetch_and_refresh)

    # ---------------- Top control bar ----------------
    def _build_top_bar(self):
        bar = ttk.Frame(self.root, padding=10)
        bar.pack(fill="x")

        ttk.Label(bar, text="Device IP:").pack(side="left")
        self.ip_var = tk.StringVar(value=self.cfg.get("device_ip", "192.168.29.201"))
        ttk.Entry(bar, textvariable=self.ip_var, width=16).pack(side="left", padx=5)

        ttk.Button(bar, text="Fetch Latest & Update",
                   command=self.fetch_and_refresh).pack(side="left", padx=5)

        ttk.Label(bar, text="  Google Drive Folder:").pack(side="left")
        self.drive_var = tk.StringVar(value=self.cfg.get("drive_sync_folder", ""))
        ttk.Entry(bar, textvariable=self.drive_var, width=35).pack(side="left", padx=5)
        ttk.Button(bar, text="Browse...",
                   command=self.browse_drive_folder).pack(side="left")

        # ── EXPORT ADDITION ── Export button in top bar ──────────────────────
        ttk.Button(
            bar,
            text="📥 Export Attendance Report",
            command=self.export_attendance_report,
        ).pack(side="left", padx=(15, 5))
        # ── END EXPORT ADDITION ──────────────────────────────────────────────

        self.status_var = tk.StringVar(value="")
        ttk.Label(self.root, textvariable=self.status_var,
                  foreground="#2980b9", padding=(10, 0)).pack(fill="x")

    def browse_drive_folder(self):
        path = filedialog.askdirectory(title="Select your local Google Drive synced folder")
        if path:
            self.drive_var.set(path)
            self.cfg["drive_sync_folder"] = path
            save_config(self.cfg)

    # ── EXPORT ADDITION ── Export handler ────────────────────────────────────
    def export_attendance_report(self):
        """
        Opens a Save-As dialog and writes a styled 3-sheet Excel report
        from the current attendance.db — mirrors Local_Attendance export.
        """
        from datetime import datetime
        default_name = f"attendance_report_{datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}.xlsx"

        output_path = filedialog.asksaveasfilename(
            title="Save Attendance Report",
            initialfile=default_name,
            defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx"), ("All Files", "*.*")],
        )
        if not output_path:
            return  # user cancelled

        try:
            standard_hours = self.cfg.get("standard_workday_hours", 9)
            export_attendance_to_excel(self.attendance_db, output_path, standard_hours)
            self.status_var.set(f"Report exported → {output_path}")
            messagebox.showinfo(
                "Export Successful",
                f"Attendance report saved to:\n{output_path}",
            )
        except FileNotFoundError as e:
            messagebox.showwarning("No Data", str(e))
        except Exception as e:
            messagebox.showerror("Export Failed", f"Could not export report:\n{e}")
    # ── END EXPORT ADDITION ──────────────────────────────────────────────────

    # ---------------- Tabs ----------------
    def _build_tabs(self):
        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill="both", expand=True, padx=10, pady=10)

        self.tab_overview = ttk.Frame(self.notebook)
        self.tab_employee = ttk.Frame(self.notebook)
        self.tab_payroll = ttk.Frame(self.notebook)

        self.notebook.add(self.tab_overview, text="📊 Dashboard Overview")
        self.notebook.add(self.tab_employee, text="🔍 Employee Drill-Down")
        self.notebook.add(self.tab_payroll, text="💰 Payroll (Per-Day Salary)")

    # ---------------- Fetch + Update pipeline ----------------
    def fetch_and_refresh(self):
        ip = self.ip_var.get().strip()
        self.cfg["device_ip"] = ip
        save_config(self.cfg)
        self.status_var.set(f"Connecting to {ip} ...")
        threading.Thread(target=self._background_fetch, args=(ip,), daemon=True).start()

    def _background_fetch(self, ip):
        try:
            records = attendance_engine.fetch_and_process(
                ip, self.cfg.get("standard_workday_hours", 9)
            )
            db_manager.update_attendance_db(self.attendance_db, records)
            daily_df, monthly_df = db_manager.load_dataframes(self.attendance_db)
            payroll_manager.compute_and_store_daily_pay(
                self.payroll_db, daily_df, self.cfg.get("default_hourly_rate", 100.0)
            )
            drive_folder = self.drive_var.get().strip()
            msg1, msg2 = "", ""
            if drive_folder:
                ok1, msg1 = db_manager.sync_file_to_drive(self.attendance_db, drive_folder)
                ok2, msg2 = db_manager.sync_file_to_drive(self.payroll_db, drive_folder)
            self.root.after(0, lambda: self._on_fetch_success(drive_folder, msg1, msg2))
        except Exception as e:
            self.root.after(0, lambda: self._on_fetch_error(str(e)))

    def _on_fetch_success(self, drive_folder, msg1, msg2):
        self.status_var.set(f"Updated successfully. {msg1} {msg2}".strip())
        self.refresh_dashboard()

    def _on_fetch_error(self, error_msg):
        self.status_var.set(f"Error: {error_msg}")
        messagebox.showerror("Fetch Failed", f"Could not fetch data:\n{error_msg}")

    # ---------------- Populate dashboard tabs ----------------
    def refresh_dashboard(self):
        daily_df, monthly_df = db_manager.load_dataframes(self.attendance_db)
        daily_pay_df, monthly_payroll_df = payroll_manager.load_payroll_data(self.payroll_db)
        self._populate_overview(monthly_df)
        self._populate_employee_tab(daily_df, monthly_df)
        self._populate_payroll_tab(daily_pay_df, monthly_payroll_df)

    def _clear_frame(self, frame):
        for widget in frame.winfo_children():
            widget.destroy()

    def _populate_overview(self, monthly_df):
        self._clear_frame(self.tab_overview)
        frame = ttk.Frame(self.tab_overview)
        frame.pack(fill="both", expand=True)
        fig1 = charts.hours_trend_figure(monthly_df)
        canvas1 = FigureCanvasTkAgg(fig1, master=frame)
        canvas1.get_tk_widget().pack(side="left", fill="both", expand=True)
        canvas1.draw()
        fig2 = charts.excess_deficit_figure(monthly_df)
        canvas2 = FigureCanvasTkAgg(fig2, master=frame)
        canvas2.get_tk_widget().pack(side="left", fill="both", expand=True)
        canvas2.draw()

    # ... (rest of _populate_employee_tab, _populate_payroll_tab unchanged)


if __name__ == "__main__":
    root = tk.Tk()
    app = AttendanceApp(root)
    root.mainloop()
